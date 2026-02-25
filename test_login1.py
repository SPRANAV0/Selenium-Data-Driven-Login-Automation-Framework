import pytest
import allure
from selenium import webdriver
from selenium.webdriver.common.by import By
from openpyxl import load_workbook

worbook = load_workbook(r"C:\Users\Admin\OneDrive\Desktop\Allure_Testing\test_data.xlsx")
sheet = worbook.active

rows = sheet.max_row

test_data = []
for i in range(2,rows+1):
    username = sheet.cell(row=i,column=1).value
    password = sheet.cell(row=i,column=2).value
    expected = sheet.cell(row=i,column=3).value
    test_data.append((username,password,expected))

@pytest.fixture
def driver():
    driver = webdriver.Chrome()
    driver.maximize_window()
    yield driver
    driver.quit()


@pytest.mark.parametrize("username,password,expected", test_data)
@allure.title("SauceDemo Login Test")

def test_login(driver,username,password,expected):

    driver.get("https://www.saucedemo.com")

    driver.find_element(By.ID,"user-name").clear()
    driver.find_element(By.ID,"user-name").send_keys(username)

    driver.find_element(By.ID,"password").clear()
    driver.find_element(By.ID,"password").send_keys(password)

    driver.find_element(By.ID,"login-button").click()

    assert "inventory" in driver.current_url
